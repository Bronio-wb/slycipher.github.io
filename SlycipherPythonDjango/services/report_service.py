"""
Servicio de generación de reportes PDF y Excel.
Extraído de SlycipherPython para su reutilización sin cursor.
"""
import os
from io import BytesIO
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.utils import simpleSplit
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from reportlab.pdfgen import canvas as pdf_canvas
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from django.conf import settings
from django.utils import timezone


# Paleta SLYCIPHER
BRAND_ORANGE = colors.HexColor('#ff6b35')
BRAND_BLUE = colors.HexColor('#263847')
BRAND_NAVY = colors.HexColor('#0f172a')
BRAND_BLUE_LIGHT = colors.HexColor('#eaf2ff')


def export_pdf(
    title,
    headers,
    rows,
    filtros=None,
    generated_by=None,
    generated_at=None,
    report_type=None,
    summary_rows=None,
    distribution_title='Distribución',
    distribution_headers=None,
    distribution_rows=None,
    include_data_section_title=False,
    data_section_title='Datos Detallados',
    details_footer_text=None,
    progress_kpis=None,
    progress_top_rows=None,
):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=26, leftMargin=26, topMargin=28, bottomMargin=34,
    )
    styles = getSampleStyleSheet()
    report_type_text = str(report_type or '').strip().lower()
    is_user_report = 'usuario' in report_type_text
    is_course_report = 'curso' in report_type_text

    meta_label_bg = colors.HexColor('#eff6ff')
    meta_border = colors.HexColor('#93c5fd')
    summary_title_bg = BRAND_BLUE
    table_header_bg = BRAND_BLUE

    title_style = styles['Title']
    title_style.textColor = BRAND_BLUE
    title_style.fontName = 'Helvetica-Bold'
    title_style.fontSize = 20

    subtitle_style = styles['Normal']
    subtitle_style.textColor = colors.HexColor('#475569')
    subtitle_style.fontSize = 9

    filter_style = styles['BodyText']
    filter_style.fontSize = 9

    detail_cell_style = styles['BodyText']
    detail_cell_style.fontName = 'Helvetica'
    detail_cell_style.fontSize = 8
    detail_cell_style.textColor = colors.HexColor('#111827')
    detail_cell_style.wordWrap = 'CJK'

    story = []

    # Cabecera con logo
    logo_path = os.path.join(settings.BASE_DIR, 'static', 'images', 'logo.png')
    header_cells = []
    if os.path.exists(logo_path):
        try:
            header_cells.append(Image(logo_path, width=34, height=34))
        except Exception:
            header_cells.append('')
    else:
        header_cells.append('')

    header_cells.append([
        Paragraph('<font color="#0f172a"><b>SLYCIPHER - Panel de Administración</b></font>', subtitle_style),
        Paragraph(title, title_style),
        Paragraph('Documento generado automáticamente desde el módulo de reportes.', subtitle_style),
    ])
    header_table = Table([header_cells], colWidths=[44, 700])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
    ]))
    story.extend([header_table, Spacer(1, 12)])

    # Metadatos
    metadata_rows = []
    if generated_by:
        metadata_rows.append(['Generado por', str(generated_by)])
    if generated_at:
        metadata_rows.append(['Fecha de generación', str(generated_at)])
    if report_type:
        metadata_rows.append(['Tipo de reporte', str(report_type)])
    if metadata_rows:
        meta_table = Table(metadata_rows, colWidths=[170, 574])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), meta_label_bg),
            ('BACKGROUND', (1, 0), (1, -1), colors.white),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTNAME', (1, 0), (1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOX', (0, 0), (-1, -1), 0.6, meta_border),
            ('INNERGRID', (0, 0), (-1, -1), 0.3, BRAND_BLUE_LIGHT),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        story.extend([meta_table, Spacer(1, 10)])

    # Progreso (especial)
    if progress_kpis:
        kpi_cells = []
        for kpi in (progress_kpis or [])[:4]:
            kpi_cells.append(Paragraph(
                f"<para align='center'><font name='Helvetica-Bold' size='24' color='#263847'>"
                f"{escape(str(kpi.get('value','0')))}</font><br/>"
                f"<font name='Helvetica' size='13' color='#1f2937'>{escape(str(kpi.get('label','Métrica')))}</font></para>",
                styles['BodyText'],
            ))
        while len(kpi_cells) < 4:
            kpi_cells.append(Paragraph('', styles['BodyText']))

        kpi_table = Table([kpi_cells], colWidths=[186, 186, 186, 186])
        kpi_table.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.8, colors.HexColor('#b7b7b7')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 18),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 14),
        ]))
        story.extend([kpi_table, Spacer(1, 20)])

        top_headers = ['Posición', 'Usuario', 'Email', 'Racha']
        top_rows = progress_top_rows or [['-', 'Sin resultados', '-', '-']]
        top_table = Table([top_headers] + top_rows, colWidths=[160, 180, 180, 160])
        top_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), table_header_bg),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 12),
            ('GRID', (0, 0), (-1, -1), 0.8, colors.HexColor('#2f2f2f')),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(top_table)
        doc.build(story)
        return buffer.getvalue()

    # Filtros
    if filtros:
        filtros_table = Table(
            [[Paragraph('<b>Filtros aplicados</b>', filter_style)],
             [Paragraph(' | '.join(filtros), filter_style)]],
            colWidths=[744],
        )
        filtros_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), BRAND_BLUE),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#fff7ed')),
            ('BOX', (0, 0), (-1, -1), 0.6, colors.HexColor('#fdba74')),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.extend([filtros_table, Spacer(1, 12)])

    # Resumen
    if summary_rows:
        sum_title = Table([[Paragraph('<b>Resumen Estadístico</b>', filter_style)]], colWidths=[744])
        sum_title.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), summary_title_bg),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.white),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        sum_table = Table(summary_rows, colWidths=[260, 484])
        sum_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
            ('INNERGRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#e5e7eb')),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.extend([sum_title, sum_table, Spacer(1, 12)])

    # Distribución
    if distribution_rows:
        dist_title = Table([[Paragraph(f'<b>{distribution_title}</b>', filter_style)]], colWidths=[744])
        dist_title.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), summary_title_bg),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.white),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        dist_headers = distribution_headers or ['TIPO', 'CANTIDAD', 'PORCENTAJE']
        dist_table = Table([dist_headers] + distribution_rows, colWidths=[300, 220, 224])
        dist_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), table_header_bg),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('GRID', (0, 0), (-1, -1), 0.45, colors.HexColor('#d1d5db')),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.extend([dist_title, dist_table, Spacer(1, 12)])

    # Tabla de datos
    if not rows:
        rows = [['Sin resultados'] + [''] * (len(headers) - 1)]

    if include_data_section_title:
        data_title = Table([[Paragraph(f'<b>{data_section_title}</b>', filter_style)]], colWidths=[744])
        data_title.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), summary_title_bg),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.white),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.extend([data_title, Spacer(1, 6)])

    wrapped_rows = [
        [Paragraph(escape(str(cell if cell is not None else '')), detail_cell_style) for cell in row]
        for row in rows
    ]
    if is_course_report and len(headers) == 5:
        col_widths = [150, 300, 90, 102, 102]
    elif is_user_report and len(headers) == 5:
        col_widths = [145, 215, 110, 95, 179]
    else:
        col_widths = [744 / max(len(headers), 1)] * len(headers)

    data_table = Table([headers] + wrapped_rows, repeatRows=1, colWidths=col_widths)
    data_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), table_header_bg),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.45, colors.HexColor('#d1d5db')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(data_table)

    if details_footer_text:
        footer_style = styles['Normal']
        footer_style.fontName = 'Helvetica-Bold'
        footer_style.fontSize = 12
        story.extend([Spacer(1, 10), Paragraph(details_footer_text, footer_style)])

    def draw_footer(canvas, document):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#6b7280'))
        ts = str(generated_at or timezone.localtime(timezone.now()).strftime('%d/%m/%Y %H:%M:%S'))
        canvas.drawString(document.leftMargin, 8, f'Reporte SLYCIPHER — {ts}')
        canvas.drawRightString(document.pagesize[0] - document.rightMargin, 8, f'Página {canvas.getPageNumber()}')
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    return buffer.getvalue()


def export_excel(sheet_name, headers, rows, filtros=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] or 'Reporte'

    title_fill = PatternFill(fill_type='solid', fgColor='263847')
    accent_fill = PatternFill(fill_type='solid', fgColor='FF6B35')
    header_fill = PatternFill(fill_type='solid', fgColor='EAF2FF')
    alt_fill = PatternFill(fill_type='solid', fgColor='F8FAFC')
    white_font = Font(color='FFFFFF', bold=True, size=14)
    section_font = Font(color='FFFFFF', bold=True, size=11)
    header_font = Font(color='1F2937', bold=True, size=11)
    body_font = Font(color='111827', size=10)
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
    )
    max_col = max(len(headers), 1)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws['A1'] = f'SLYCIPHER - {sheet_name}'
    ws['A1'].fill = title_fill
    ws['A1'].font = white_font
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 24

    current_row = 2
    if filtros:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_col)
        ws.cell(row=current_row, column=1, value='Filtros: ' + ' | '.join(filtros))
        ws.cell(row=current_row, column=1).fill = accent_fill
        ws.cell(row=current_row, column=1).font = section_font
        ws.cell(row=current_row, column=1).alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[current_row].height = 22
        current_row += 2

    header_row = current_row
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=ci, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    data_start = header_row + 1
    for ri, row in enumerate(rows):
        row_idx = data_start + ri
        for ci, val in enumerate(list(row), 1):
            ws.cell(row=row_idx, column=ci, value=val)

    for ri in range(data_start, ws.max_row + 1):
        use_alt = (ri - data_start) % 2 == 1
        for ci in range(1, len(headers) + 1):
            cell = ws.cell(row=ri, column=ci)
            cell.font = body_font
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = thin_border
            if use_alt:
                cell.fill = alt_fill

    ws.auto_filter.ref = f'A{header_row}:{ws.cell(row=header_row, column=len(headers)).coordinate}'
    ws.freeze_panes = f'A{data_start}'

    for ci in range(1, len(headers) + 1):
        col_letter = ws.cell(row=header_row, column=ci).column_letter
        max_len = max((len(str(ws.cell(row=ri, column=ci).value or '')) for ri in range(1, ws.max_row + 1)), default=10)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_student_certificate_pdf(
    student_name,
    completed_courses,
    total_lessons_completed,
    total_challenges_completed,
    generated_by=None,
    generated_at=None,
    certificate_kind='general',
    course_name=None,
    certificate_code=None,
    ceremonial_date=None,
    validation_url=None,
):
    buffer = BytesIO()
    page_width, page_height = A4
    pdf = pdf_canvas.Canvas(buffer, pagesize=A4)

    generated_at = generated_at or timezone.localtime().strftime('%d/%m/%Y %H:%M:%S')

    pdf.setFillColor(colors.HexColor('#f8fafc'))
    pdf.rect(0, 0, page_width, page_height, stroke=0, fill=1)

    pdf.setFillColor(BRAND_BLUE)
    pdf.rect(0, page_height - 86, page_width, 86, stroke=0, fill=1)

    pdf.setFillColor(colors.HexColor('#fff7ed'))
    pdf.rect(34, 34, page_width - 68, page_height - 120, stroke=0, fill=1)

    pdf.setStrokeColor(BRAND_NAVY)
    pdf.setLineWidth(1.4)
    pdf.rect(28, 28, page_width - 56, page_height - 56, stroke=1, fill=0)

    pdf.setStrokeColor(BRAND_ORANGE)
    pdf.setLineWidth(3)
    pdf.rect(40, 40, page_width - 80, page_height - 80, stroke=1, fill=0)

    logo_path = os.path.join(settings.BASE_DIR, 'static', 'images', 'logo.png')
    if os.path.exists(logo_path):
        try:
            pdf.saveState()
            pdf.setFillAlpha(0.08)
            pdf.drawImage(
                logo_path,
                (page_width / 2) - 140,
                (page_height / 2) - 140,
                width=280,
                height=280,
                preserveAspectRatio=True,
                mask='auto',
            )
            pdf.restoreState()
        except Exception:
            pass

    if os.path.exists(logo_path):
        try:
            # Logo alineado a la izquierda del bloque de texto del encabezado.
            pdf.drawImage(logo_path, 54, page_height - 70, width=44, height=44, preserveAspectRatio=True, mask='auto')
        except Exception:
            pass

    pdf.setFillColor(colors.white)
    pdf.setFont('Helvetica-Bold', 18)
    pdf.drawString(112, page_height - 56, 'SLYCIPHER')
    pdf.setFont('Helvetica', 10)
    pdf.drawString(112, page_height - 72, 'Plataforma de formacion y logros academicos')

    certificate_title = 'CERTIFICADO DE LOGRO' if certificate_kind == 'general' else 'CERTIFICADO DE CURSO'
    certificate_subtitle = (
        'SLYCIPHER reconoce el esfuerzo academico del estudiante'
        if certificate_kind == 'general'
        else 'SLYCIPHER certifica la aprobacion satisfactoria del curso'
    )

    pdf.setFillColor(BRAND_BLUE)
    pdf.setFont('Helvetica-Bold', 30)
    pdf.drawCentredString(page_width / 2, page_height - 145, certificate_title)

    pdf.setFillColor(colors.HexColor('#475569'))
    pdf.setFont('Helvetica', 13)
    pdf.drawCentredString(page_width / 2, page_height - 173, certificate_subtitle)

    pdf.setStrokeColor(colors.HexColor('#cbd5e1'))
    pdf.line(135, page_height - 192, page_width - 135, page_height - 192)

    pdf.setFillColor(BRAND_ORANGE)
    pdf.setFont('Helvetica-Bold', 28)
    pdf.drawCentredString(page_width / 2, page_height - 240, str(student_name or 'Estudiante'))

    pdf.setFillColor(colors.black)
    body_font_name = 'Helvetica'
    body_font_size = 14
    pdf.setFont(body_font_name, body_font_size)
    if certificate_kind == 'general':
        completion_text = 'por completar satisfactoriamente multiples cursos, lecciones y desafios en la plataforma.'
    else:
        completion_text = f'por completar satisfactoriamente el curso {course_name or "asignado"}.'

    # Ajuste de ancho para que el texto nunca salga del marco interior.
    text_max_width = page_width - 2 * 58
    wrapped_lines = simpleSplit(completion_text, body_font_name, body_font_size, text_max_width)
    y_base = page_height - 278
    line_step = 18
    for index, line in enumerate(wrapped_lines[:3]):
        pdf.drawCentredString(page_width / 2, y_base - (index * line_step), line)

    metrics_top = page_height - 382
    box_width = 150
    start_x = 60
    gap = 24
    metric_items = [
        ('Cursos completados', len(completed_courses or [])),
        ('Lecciones completadas', total_lessons_completed),
        ('Desafios correctos', total_challenges_completed),
    ]
    for index, (label, value) in enumerate(metric_items):
        x = start_x + (box_width + gap) * index
        pdf.setStrokeColor(colors.HexColor('#cbd5e1'))
        pdf.setFillColor(colors.HexColor('#f8fafc'))
        pdf.roundRect(x, metrics_top, box_width, 72, 10, stroke=1, fill=1)
        pdf.setFillColor(BRAND_BLUE)
        pdf.setFont('Helvetica-Bold', 22)
        pdf.drawCentredString(x + (box_width / 2), metrics_top + 40, str(value))
        pdf.setFillColor(colors.HexColor('#475569'))
        pdf.setFont('Helvetica', 10)
        pdf.drawCentredString(x + (box_width / 2), metrics_top + 18, label)

    pdf.setFillColor(BRAND_BLUE)
    pdf.setFont('Helvetica-Bold', 14)
    pdf.drawString(58, page_height - 478, 'Cursos acreditados:' if certificate_kind == 'general' else 'Curso certificado:')

    pdf.setFont('Helvetica', 12)
    y = page_height - 503
    for course in (completed_courses or [])[:8]:
        pdf.drawString(75, y, f'- {course}')
        y -= 20

    if certificate_code:
        pdf.setFillColor(colors.HexColor('#0f172a'))
        pdf.roundRect(56, 146, 180, 30, 8, stroke=0, fill=1)
        pdf.setFillColor(colors.white)
        pdf.setFont('Helvetica-Bold', 11)
        pdf.drawString(70, 157, f'Codigo: {certificate_code}')

    if validation_url:
        pdf.setFillColor(colors.HexColor('#475569'))
        pdf.setFont('Helvetica', 8)
        pdf.drawString(58, 136, 'Validacion publica:')
        pdf.setFillColor(BRAND_BLUE)
        pdf.drawString(132, 136, validation_url[:74])

    pdf.setStrokeColor(colors.HexColor('#94a3b8'))
    pdf.line(58, 120, 220, 120)
    pdf.line(page_width - 220, 120, page_width - 58, 120)
    pdf.setFont('Helvetica', 11)
    pdf.setFillColor(colors.HexColor('#475569'))
    pdf.drawString(58, 104, 'Firma autorizada')
    pdf.drawString(page_width - 210, 104, f'Fecha: {ceremonial_date or generated_at}')

    pdf.setFont('Helvetica-Bold', 12)
    pdf.setFillColor(BRAND_BLUE)
    pdf.drawString(58, 88, generated_by or 'Administrador SLYCIPHER')
    pdf.setFont('Helvetica', 10)
    pdf.setFillColor(colors.HexColor('#475569'))
    pdf.drawString(58, 74, 'Administrador SLYCIPHER')

    pdf.setFont('Helvetica', 9)
    pdf.setFillColor(colors.HexColor('#64748b'))
    pdf.drawCentredString(page_width / 2, 62, 'Documento generado automaticamente desde el panel de administracion SLYCIPHER')

    pdf.showPage()
    pdf.save()
    return buffer.getvalue()
